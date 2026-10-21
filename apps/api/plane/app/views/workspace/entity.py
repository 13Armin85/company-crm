"""Metadata administration and a schema-driven record API, scoped to a workspace."""

from copy import deepcopy
from io import BytesIO
from pathlib import PurePath

from django.conf import settings
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, FileResponse, Http404
from django.shortcuts import get_object_or_404
from rest_framework.authentication import SessionAuthentication
from rest_framework.exceptions import ValidationError, PermissionDenied, APIException
from rest_framework.response import Response

from plane.authentication.session import BaseSessionAuthentication
from plane.app.views.base import BaseAPIView
from plane.app.permissions.crm import WorkspaceScopePermission
from plane.app.services.access_control import has_permission
from plane.app.services.entity_schema import (
    FIELD_KINDS,
    field_dict,
    bindings,
    schema_fields,
    present_schema,
    validate_layout,
    can_field,
    role_ids,
)
from plane.app.services.entity_records import (
    ADAPTERS,
    permission,
    require_action,
    native_queryset,
    record_queryset,
    get_record,
    all_values,
    record_payload,
    save_record,
    delete_record,
    references_to,
    emit,
    emit_definition,
)
from plane.app.services.entity_config import export_config, import_config, TEMPLATES, template_config
from plane.app.serializers.entity import (
    EntityTypeSerializer,
    LookupCategorySerializer,
    LookupValueSerializer,
    FieldDefinitionSerializer,
)
from plane.db.models import (
    Workspace,
    EntityType,
    EntityRecord,
    EntityEvent,
    LookupCategory,
    LookupValue,
    FieldDefinition,
    FormLayout,
    FormVersion,
    OrganizationRole,
    OrganizationUnit,
    User,
    FileAsset,
    State,
)


DEFINITIONS = {
    "entities": (EntityType, EntityTypeSerializer, "EntityType"),
    "lookups": (LookupCategory, LookupCategorySerializer, "Lookup"),
    "values": (LookupValue, LookupValueSerializer, "Lookup"),
    "fields": (FieldDefinition, FieldDefinitionSerializer, "Field"),
}


def conflict(detail):
    error = APIException(detail)
    error.status_code = 409
    raise error


def page(request):
    try:
        return max(0, int(request.query_params.get("page", 1)) - 1), min(
            100, max(1, int(request.query_params.get("page_size", 30)))
        )
    except (ValueError, TypeError):
        raise ValidationError("شماره یا اندازه صفحه معتبر نیست.")


class EngineAPIView(BaseAPIView):
    permission_classes = [WorkspaceScopePermission]

    def workspace(self, slug):
        return get_object_or_404(Workspace, slug=slug)

    def entity(self, slug, code):
        return get_object_or_404(
            EntityType.objects.select_related("workspace"), workspace__slug=slug, code=code, parent__isnull=True
        )

    def require(self, request, workspace, code):
        if not has_permission(request.user, workspace, code, request=request):
            raise PermissionDenied("مجوز انجام این عملیات را ندارید.")

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if request.method not in {"GET", "HEAD", "OPTIONS"} and isinstance(
            request.successful_authenticator, BaseSessionAuthentication
        ):
            SessionAuthentication().enforce_csrf(request)

    def finalize_response(self, request, response, *args, **kwargs):
        response = super().finalize_response(request, response, *args, **kwargs)
        response["Cache-Control"] = "private, no-store"
        return response


def definition_audit(request, workspace, action, row, before=None, after=None):
    emit_definition(workspace, request.user, action, row, before, after)


def field_referenced(workspace, code):
    from plane.db.models import SurfaceConfiguration
    if any(b.get("field") == code for definition in SurfaceConfiguration.objects.filter(workspace=workspace).values_list("definition", flat=True) for b in definition.get("fields", [])):
        return True
    def mentions(value):
        if isinstance(value, dict):
            return (
                value.get("field") == code
                or value.get("depends_on") == code
                or any(mentions(child) for child in value.values())
            )
        return isinstance(value, list) and any(mentions(child) for child in value)

    if any(
        mentions(config)
        for config in FieldDefinition.objects.filter(workspace=workspace)
        .exclude(code=code)
        .values_list("config", flat=True)
    ):
        return True
    return any(
        b.get("field") == code
        for form in FormLayout.objects.filter(entity__workspace=workspace)
        for layout in [form.draft, *form.versions.values_list("schema", flat=True)]
        for b in bindings(layout)
    )


def lookup_used(value):
    fields = FieldDefinition.objects.filter(lookup=value.category)
    for field in fields:
        key = value.code if field.config.get("lookup_codes") else str(value.id)

        def selected(raw):
            return raw == key or isinstance(raw, list) and key in raw

        def condition_uses(rule):
            return isinstance(rule, dict) and (
                rule.get("field") == field.code
                and selected(rule.get("value"))
                or any(condition_uses(child) for child in rule.get("rules", []))
            )

        if selected(field.config.get("default")):
            return True
        for config in FieldDefinition.objects.filter(workspace=value.category.workspace).values_list(
            "config", flat=True
        ):
            if any(condition_uses(config.get(name)) for name in ("required_if", "visible_if")):
                return True
        for form in FormLayout.objects.filter(entity__workspace=value.category.workspace):
            for layout in [form.draft, *form.versions.values_list("schema", flat=True)]:
                for binding in bindings(layout):
                    if (
                        binding.get("field") == field.code
                        and selected(binding.get("default"))
                        or any(condition_uses(binding.get(name)) for name in ("required_if", "visible_if"))
                    ):
                        return True
        if (
            EntityRecord.all_objects.filter(workspace=value.category.workspace)
            .filter(Q(**{f"data__{field.code}": key}) | Q(**{f"data__{field.code}__contains": [key]}))
            .exists()
        ):
            return True
        for entity in EntityType.objects.filter(workspace=value.category.workspace, parent__isnull=True).exclude(
            adapter=""
        ):
            _, schema = schema_fields(entity)
            source = schema.get(field.code, {}).get("binding", {}).get("source")
            if (
                source
                and ADAPTERS[entity.adapter].all_objects.filter(workspace=entity.workspace, **{source: key}).exists()
            ):
                return True
    return False


class EngineCatalogEndpoint(EngineAPIView):
    def get(self, request, slug):
        workspace = self.workspace(slug)
        entities = EntityType.objects.filter(workspace=workspace).select_related("parent")
        rows = [
            dict(
                EntityTypeSerializer(e, context={"workspace": workspace}).data,
                can_create=permission(e.parent or e, request.user, "create", request),
            )
            for e in entities
            if has_permission(request.user, workspace, "EntityType.View", request=request)
            or permission(e.parent or e, request.user, "view", request)
        ]
        choices = {"priority": [], "task_status": []}
        roles = role_ids(request.user, workspace)
        for entity in entities:
            if entity.parent_id or not permission(entity, request.user, "view", request):
                continue
            _, fields = schema_fields(entity)
            for code in choices:
                field = fields.get(code)
                if field and field.get("lookup") and can_field(field, roles) and not field["config"].get("hidden"):
                    choices[code] = [
                        {
                            "id": option.code,
                            "title": option.title,
                            "color": option.color,
                            "disabled": not option.is_active or not option.category.is_active,
                        }
                        for option in LookupValue.objects.filter(
                            category__workspace=workspace, category__code=field["lookup"]
                        ).select_related("category")
                    ]
        return Response(
            {
                "entities": rows,
                "field_kinds": [{"code": key, "title": title} for key, title in FIELD_KINDS.items()],
                "templates": [{"code": key, "title": value["title"]} for key, value in TEMPLATES.items()],
                "choice_labels": choices,
            }
        )


class EngineDefinitionEndpoint(EngineAPIView):
    def scope(self, request, slug, kind, action, identifier=None):
        if kind not in DEFINITIONS:
            raise ValidationError("نوع تعریف معتبر نیست.")
        model, serializer, prefix = DEFINITIONS[kind]
        workspace = self.workspace(slug)
        self.require(request, workspace, f"{prefix}.{action}")
        if action != "View":
            Workspace.objects.select_for_update().get(id=workspace.id)
        rows = model.objects.filter(
            **({"category__workspace": workspace} if kind == "values" else {"workspace": workspace})
        )
        row = (
            get_object_or_404(rows.select_for_update() if action != "View" else rows, id=identifier)
            if identifier
            else None
        )
        return workspace, rows, row, serializer

    @transaction.atomic
    def get(self, request, slug, kind, identifier=None):
        workspace, rows, row, serializer = self.scope(request, slug, kind, "View", identifier)
        if kind == "values" and request.query_params.get("category"):
            rows = rows.filter(category_id=request.query_params["category"])
        return Response(serializer(row if row else rows, many=not bool(row), context={"workspace": workspace}).data)

    @transaction.atomic
    def post(self, request, slug, kind, identifier=None):
        workspace, _, _, serializer_type = self.scope(request, slug, kind, "Create")
        Workspace.objects.select_for_update().get(id=workspace.id)
        serializer = serializer_type(data=request.data, context={"workspace": workspace})
        serializer.is_valid(raise_exception=True)
        if kind == "values" and serializer.validated_data.get("is_default"):
            LookupValue.objects.filter(category=serializer.validated_data["category"]).update(is_default=False)
        row = serializer.save(**({} if kind == "values" else {"workspace": workspace}))
        definition_audit(request, workspace, f"definition.{kind}.created", row, after=dict(serializer.data))
        return Response(serializer.data, status=201)

    @transaction.atomic
    def patch(self, request, slug, kind, identifier):
        workspace, _, row, serializer_type = self.scope(request, slug, kind, "Edit", identifier)
        Workspace.objects.select_for_update().get(id=workspace.id)
        before = dict(serializer_type(row, context={"workspace": workspace}).data)
        serializer = serializer_type(row, data=request.data, partial=True, context={"workspace": workspace})
        serializer.is_valid(raise_exception=True)
        if kind == "values" and serializer.validated_data.get("is_default"):
            LookupValue.objects.filter(category=row.category).exclude(id=row.id).update(is_default=False)
        serializer.save()
        definition_audit(request, workspace, f"definition.{kind}.updated", row, before, dict(serializer.data))
        return Response(serializer.data)

    @transaction.atomic
    def delete(self, request, slug, kind, identifier):
        workspace, _, row, serializer_type = self.scope(request, slug, kind, "Delete", identifier)
        Workspace.objects.select_for_update().get(id=workspace.id)
        if kind == "entities" and (
            row.adapter
            or EntityRecord.all_objects.filter(Q(entity=row) | Q(subtype=row)).exists()
            or row.subtypes.exists()
            or row.relation_fields.exists()
        ):
            raise ValidationError("موجودیت سیستمی یا استفاده‌شده حذف نمی‌شود؛ آن را غیرفعال کنید.")
        if kind == "lookups" and (
            row.is_system
            or row.fields.exists()
            or any(lookup_used(value) for value in row.values.all())
            or LookupValue.objects.filter(parent__category=row).exclude(category=row).exists()
        ):
            raise ValidationError("لیست سیستمی یا دارای وابستگی حذف نمی‌شود؛ آن را غیرفعال کنید.")
        if kind == "values" and (lookup_used(row) or row.children.exists()):
            raise ValidationError("این مقدار استفاده شده یا والد مقادیر دیگر است؛ آن را غیرفعال کنید.")
        if kind == "fields" and (row.used_at or field_referenced(workspace, row.code)):
            raise ValidationError("فیلد استفاده‌شده یا متصل به فرم حذف نمی‌شود؛ آن را غیرفعال کنید.")
        before = dict(serializer_type(row, context={"workspace": workspace}).data)
        definition_audit(request, workspace, f"definition.{kind}.deleted", row, before)
        if kind == "lookups":
            # Remove unused child values before their parent; external dependencies were checked above.
            row.values.update(parent=None)
            row.values.all().delete(soft=False)
        if kind == "entities":
            from plane.app.services.entity_issue_types import native_type_for_subtype
            from plane.db.models import Issue

            native_type = native_type_for_subtype(row, create=False)
            if native_type:
                if Issue.all_objects.filter(workspace=workspace, type=native_type).exists():
                    raise ValidationError("این زیرنوع در تیکت‌های فعلی یا تاریخی استفاده شده است؛ آن را غیرفعال کنید.")
                native_type.delete(soft=False)
            form = FormLayout.objects.filter(entity=row).first()
            if form:
                form.published = None
                form.save(update_fields=["published", "updated_at"])
                form.versions.all().delete(soft=False)
                form.delete(soft=False)
        row.delete(soft=False)
        return Response(status=204)


class EngineFormEndpoint(EngineAPIView):
    def form(self, request, slug, code, action):
        workspace = self.workspace(slug)
        self.require(request, workspace, action)
        entity = get_object_or_404(EntityType, workspace=workspace, code=code)
        return entity

    def get(self, request, slug, code):
        entity = self.form(request, slug, code, "FormLayout.Edit")
        form = FormLayout.objects.filter(entity=entity).select_related("published").first()
        return Response(
            {
                "draft": form.draft if form else {"tabs": [], "sections": []},
                "revision": form.revision if form else 0,
                "published": form.published.schema if form and form.published_id else None,
                "versions": list(form.versions.values("id", "number", "created_at", "created_by_id", "change_note"))
                if form
                else [],
            }
        )

    @transaction.atomic
    def put(self, request, slug, code):
        entity = self.form(request, slug, code, "FormLayout.Edit")
        EntityType.objects.select_for_update().get(id=entity.id)
        form = FormLayout.objects.select_for_update().filter(entity=entity).first()
        if request.data.get("revision") != (form.revision if form else 0):
            conflict("فرم توسط مدیر دیگری تغییر کرده است؛ دوباره بارگذاری کنید.")
        draft = validate_layout(entity, request.data.get("draft"))
        before = deepcopy(form.draft) if form else None
        if not form:
            form = FormLayout(entity=entity, draft=draft)
        else:
            form.draft, form.revision = draft, form.revision + 1
        form.save()
        emit(
            entity.parent or entity,
            request.user,
            "definition.form_draft_saved",
            before=before,
            after=draft,
            metadata={"form_entity": entity.code},
        )
        return Response({"revision": form.revision})

    @transaction.atomic
    def post(self, request, slug, code):
        entity = self.form(request, slug, code, "FormLayout.Publish")
        EntityType.objects.select_for_update().get(id=entity.id)
        form = get_object_or_404(FormLayout.objects.select_for_update(), entity=entity)
        if request.data.get("revision") != form.revision:
            conflict("فرم تغییر کرده است؛ نسخه جدید را بررسی کنید و دوباره انتشار دهید.")
        source_version = request.data.get("rollback_version")
        snapshot = get_object_or_404(form.versions, number=source_version).schema if source_version else form.draft
        validate_layout(entity, snapshot)
        number = (form.versions.order_by("-number").values_list("number", flat=True).first() or 0) + 1
        version = FormVersion.objects.create(
            layout=form,
            number=number,
            schema=deepcopy(snapshot),
            change_note=str(request.data.get("change_note", ""))[:2000],
            created_by_id=request.user.id,
        )
        before = form.published.schema if form.published_id else None
        form.published, form.draft, form.revision = version, deepcopy(snapshot), form.revision + 1
        form.save()
        emit(
            entity.parent or entity,
            request.user,
            "definition.form_rolled_back" if source_version else "definition.form_published",
            before=before,
            after=snapshot,
            metadata={"version": number, "form_entity": entity.code},
        )
        return Response({"revision": form.revision, "version": number})


class EngineSchemaEndpoint(EngineAPIView):
    def get(self, request, slug, code):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "view", request)
        subtype = (
            get_object_or_404(EntityType, parent=entity, code=request.query_params["subtype"])
            if request.query_params.get("subtype")
            else None
        )
        values = {}
        row, native = None, None
        if request.query_params.get("record"):
            row, native = get_record(entity, request.user, request.query_params["record"], request)
            subtype = row.subtype if row else None
            values = all_values(entity, row, native)
        schema = present_schema(entity, request.user, subtype, values)
        if row is not None or native is not None:
            from plane.app.services.entity_records import record_edit_actions, field_edit_action

            actions = record_edit_actions(entity, request.user, native, request)
            for field in schema["fields"]:
                field["config"]["readonly"] = bool(
                    field["config"].get("readonly")
                    or field["config"].get("editable_after_create") is False
                    or not actions[field_edit_action(entity, field)]
                )
        # Return evaluated defaults without adding current-user values to the definition itself.
        from plane.app.services.entity_records import default_value

        request.user._engine_workspace = entity.workspace
        schema["defaults"] = {
            field["code"]: default_value(field, request.user)
            for field in schema["fields"]
            if field["is_active"] and not field["config"].get("readonly")
        }
        return Response(schema)


class EngineRecordsEndpoint(EngineAPIView):
    def get(self, request, slug, code, identifier=None):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "view", request)
        if identifier:
            row, native = get_record(entity, request.user, identifier, request)
            return Response(record_payload(entity, request.user, row, native, request))
        offset, size = page(request)
        query = request.query_params.get("search", "")[:200]
        _, fields = schema_fields(entity)
        roles = role_ids(request.user, entity.workspace)
        searchable = {
            key: field
            for key, field in fields.items()
            if field["searchable"] and can_field(field, roles) and not field["config"].get("hidden")
        }
        if entity.adapter:
            rows = native_queryset(entity, request.user, request)
            if query:
                predicate = Q(pk__in=[])
                for key, field in searchable.items():
                    source = field["binding"].get("source")
                    if source and field["kind"] in {"short_text", "long_text", "rich_text"}:
                        predicate |= Q(**{f"{source}__icontains": query})
                    elif not source:
                        predicate |= Q(
                            id__in=EntityRecord.objects.filter(
                                entity=entity, **{f"data__{key}__icontains": query}
                            ).values("source_id")
                        )
                rows = rows.filter(predicate)
            count = rows.count()
            natives = list(rows.order_by("-created_at", "id")[offset * size : (offset + 1) * size])
            sidecars = {
                r.source_id: r
                for r in EntityRecord.objects.filter(
                    entity=entity, source_id__in=[n.id for n in natives]
                ).select_related("subtype")
            }
            results = [record_payload(entity, request.user, sidecars.get(n.id), n, request) for n in natives]
        else:
            rows = record_queryset(entity, request.user, request).select_related("subtype")
            if query:
                predicate = Q(pk__in=[])
                for key in searchable:
                    predicate |= Q(**{f"data__{key}__icontains": query})
                rows = rows.filter(predicate)
            count = rows.count()
            results = [
                record_payload(entity, request.user, r, request=request)
                for r in rows.order_by("-created_at", "id")[offset * size : (offset + 1) * size]
            ]
        return Response({"results": results, "count": count, "page": offset + 1, "page_size": size})

    def post(self, request, slug, code):
        return Response(save_record(self.entity(slug, code), request.user, request.data, request), status=201)

    def patch(self, request, slug, code, identifier):
        entity = self.entity(slug, code)
        if "is_active" in request.data:
            if type(request.data["is_active"]) is not bool:
                raise ValidationError("وضعیت فعالیت باید بله/خیر باشد.")
            delete_record(
                entity,
                request.user,
                identifier,
                request,
                deactivate=not bool(request.data["is_active"]),
                confirm=request.data.get("confirm_relations") is True,
            )
            return Response({"is_active": bool(request.data["is_active"])})
        return Response(save_record(entity, request.user, request.data, request, identifier))

    def delete(self, request, slug, code, identifier):
        delete_record(
            self.entity(slug, code),
            request.user,
            identifier,
            request,
            confirm=request.query_params.get("confirm_relations") == "true",
        )
        return Response(status=204)


class EngineOptionsEndpoint(EngineAPIView):
    def get(self, request, slug, code, field_code):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "view", request)
        subtype = EntityType.objects.filter(parent=entity, code=request.query_params.get("subtype")).first()
        _, fields = schema_fields(entity, subtype)
        field = fields.get(field_code)
        if request.query_params.get("library") == "true":
            self.require(request, entity.workspace, "Field.View")
            field = field_dict(
                get_object_or_404(
                    FieldDefinition.objects.select_related("lookup", "relation_entity"),
                    workspace=entity.workspace,
                    code=field_code,
                )
            )
            field["binding"] = {}
        if not field or not can_field(field, role_ids(request.user, entity.workspace)) or field["config"].get("hidden"):
            raise PermissionDenied("مجوز مشاهده این فیلد را ندارید.")
        query, (offset, size) = request.query_params.get("search", "")[:200], page(request)
        selected = [item for item in request.query_params.get("selected", "").split(",") if item][:100]
        resolving = request.query_params.get("resolve") == "true" and bool(selected)
        config, kind = field["config"], field["kind"]
        if field.get("lookup"):
            rows = LookupValue.objects.filter(category__workspace=entity.workspace, category__code=field["lookup"])
            rows = rows.filter(
                Q(is_active=True, category__is_active=True)
                | Q(**({"code__in": selected} if config.get("lookup_codes") else {"id__in": selected}))
                if selected
                else Q(is_active=True, category__is_active=True)
            )
            if resolving:
                rows = rows.filter(**({"code__in": selected} if config.get("lookup_codes") else {"id__in": selected}))
            if request.query_params.get("parent"):
                parent = request.query_params["parent"]
                parent_field = FieldDefinition.objects.filter(
                    workspace=entity.workspace, code=config.get("depends_on")
                ).first()
                if parent_field and parent_field.config.get("lookup_codes"):
                    rows = rows.filter(parent__category_id=parent_field.lookup_id, parent__code=parent)
                else:
                    rows = rows.filter(parent_id=parent)
            elif config.get("depends_on") and not resolving:
                rows = rows.none()
            rows = rows.filter(title__icontains=query)
            count = rows.count()
            result = [
                {
                    "id": r.code if config.get("lookup_codes") else str(r.id),
                    "title": r.title,
                    "color": r.color,
                    "disabled": not r.is_active,
                }
                for r in rows[offset * size : (offset + 1) * size]
            ]
        elif config.get("provider") == "project_states":
            from plane.app.services.project_visibility import visible_projects
            from plane.db.models import Project

            project = request.query_params.get("parent")
            allowed = visible_projects(Project.objects.filter(workspace=entity.workspace), request.user, slug)
            rows = (
                State.objects.filter(project_id=project, project_id__in=allowed.values("id"), name__icontains=query)
                if project
                else State.objects.none()
            )
            if resolving:
                rows = rows.filter(id__in=selected)
            count, result = (
                rows.count(),
                [
                    {"id": str(r.id), "title": r.name, "color": r.color}
                    for r in rows[offset * size : (offset + 1) * size]
                ],
            )
        elif kind == "relation":
            target = get_object_or_404(EntityType, workspace=entity.workspace, code=field["relation_entity"])
            require_action(target, request.user, "view", request)
            rows = (
                native_queryset(target, request.user, request)
                if target.adapter
                else record_queryset(target, request.user, request)
            )
            if resolving:
                rows = rows.filter(id__in=selected)
            title_allowed = present_schema(target, request.user)["fields"]
            if target.title_field not in {
                f["code"] for f in title_allowed if not f["config"].get("hidden") and not f["config"].get("visible_if")
            }:
                rows = rows.none()
            if target.adapter:
                rows = rows.filter(**{("reason" if target.adapter == "leave" else "name") + "__icontains": query})
            else:
                rows = rows.filter(**{f"data__{target.title_field}__icontains": query})
            count, result = rows.count(), []
            for r in rows.order_by("-created_at", "id")[offset * size : (offset + 1) * size]:
                payload = (
                    record_payload(target, request.user, native=r, request=request)
                    if target.adapter
                    else record_payload(target, request.user, r, request=request)
                )
                result.append({"id": payload["id"], "title": payload["title"]})
        else:
            model = {"user": User, "unit": OrganizationUnit, "role": OrganizationRole}.get(kind)
            if not model:
                return Response({"results": [], "count": 0})
            rows = (
                User.objects.filter(
                    member_workspace__workspace=entity.workspace,
                    member_workspace__is_active=True,
                    member_workspace__deleted_at__isnull=True,
                    is_active=True,
                ).filter(
                    Q(display_name__icontains=query) | Q(first_name__icontains=query) | Q(last_name__icontains=query)
                )
                if kind == "user"
                else model.objects.filter(
                    workspace=entity.workspace,
                    is_active=True,
                    **{("name" if kind == "role" else "title") + "__icontains": query},
                )
            )
            if resolving:
                rows = rows.filter(id__in=selected)
            count = rows.distinct().count()
            result = [
                {
                    "id": str(r.id),
                    "title": (r.display_name or r.first_name or r.email)
                    if kind == "user"
                    else r.name
                    if kind == "role"
                    else r.title,
                }
                for r in rows.distinct().order_by("id")[offset * size : (offset + 1) * size]
            ]
        return Response({"results": result, "count": count, "page": offset + 1})


class EngineRelationsEndpoint(EngineAPIView):
    def get(self, request, slug, code, identifier):
        entity = self.entity(slug, code)
        get_record(entity, request.user, identifier, request)
        results = []
        for source, field, source_id in references_to(entity, identifier):
            if not permission(source, request.user, "view", request):
                continue
            try:
                row, native = get_record(source, request.user, source_id, request)
            except Exception as exc:
                from django.http import Http404

                if isinstance(exc, Http404):
                    continue
                raise
            _, fields = schema_fields(source, row.subtype if row else None)
            definition = fields.get(field.code)
            if (
                not definition
                or not can_field(definition, role_ids(request.user, entity.workspace))
                or definition["config"].get("hidden")
            ):
                continue
            payload = record_payload(source, request.user, row, native, request)
            if field.code not in payload["values"]:
                continue
            results.append(payload)
            if len(results) >= 100:
                break
        return Response({"results": results})


class EngineConfigEndpoint(EngineAPIView):
    def get(self, request, slug):
        workspace = self.workspace(slug)
        self.require(request, workspace, "Config.Export")
        return Response(export_config(workspace))

    def post(self, request, slug):
        workspace = self.workspace(slug)
        self.require(request, workspace, "Config.Import")
        self.require(request, workspace, "FormLayout.Publish")
        upload = request.FILES.get("file")
        if upload:
            import json

            if upload.size > 32 * 1024 * 1024:
                raise ValidationError("فایل پیکربندی حداکثر سی‌ودو مگابایت است.")
            try:
                return Response(import_config(workspace, request.user, json.load(upload)))
            except (ValueError, UnicodeDecodeError) as exc:
                raise ValidationError("فایل پیکربندی JSON معتبر نیست.") from exc
        config = (
            template_config(workspace, request.data["template"])
            if request.data.get("template")
            else request.data.get("config")
        )
        return Response(import_config(workspace, request.user, config))


class EngineLookupExcelEndpoint(EngineAPIView):
    def get(self, request, slug, identifier):
        workspace = self.workspace(slug)
        self.require(request, workspace, "Lookup.View")
        category = get_object_or_404(LookupCategory, workspace=workspace, id=identifier)
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(
            ["code", "title", "color", "icon", "position", "is_default", "is_active", "parent_category", "parent_code"]
        )
        for row in category.values.select_related("parent", "parent__category"):
            data = [
                row.code,
                row.title,
                row.color,
                row.icon,
                row.position,
                row.is_default,
                row.is_active,
                row.parent.category.code if row.parent_id else "",
                row.parent.code if row.parent_id else "",
            ]
            sheet.append(
                [
                    ("'" + value) if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
                    for value in data
                ]
            )
        stream = BytesIO()
        book.save(stream)
        response = HttpResponse(
            stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{category.code}.xlsx"'
        return response

    @transaction.atomic
    def post(self, request, slug, identifier):
        workspace = self.workspace(slug)
        self.require(request, workspace, "Lookup.Create")
        self.require(request, workspace, "Lookup.Edit")
        category = get_object_or_404(LookupCategory.objects.select_for_update(), workspace=workspace, id=identifier)
        upload = request.FILES.get("file")
        if not upload or upload.size > 5 * 1024 * 1024:
            raise ValidationError("یک فایل Excel با حجم حداکثر پنج مگابایت انتخاب کنید.")
        from openpyxl import load_workbook

        try:
            import zipfile

            with zipfile.ZipFile(upload) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError()
            upload.seek(0)
            book = load_workbook(upload, read_only=True, data_only=False)
            rows = iter(book.active.iter_rows(values_only=True))
            headers = list(next(rows))
            if headers[:2] != ["code", "title"]:
                raise ValueError()
            items = []
            for index, row in enumerate(rows, 2):
                if index > 10002:
                    raise ValueError()
                if not any(value is not None for value in row):
                    continue
                item = dict(zip(headers, row))
                item = {key: value for key, value in item.items() if value is not None}
                if any(isinstance(value, str) and value.startswith("=") for value in item.values()):
                    raise ValidationError({str(index): "فرمول در فایل ورودی پذیرفته نیست."})
                item["category"] = category.code
                item["_source_row"] = index
                parent_code = item.pop("parent_code", None)
                parent_category = item.pop("parent_category", None)
                item["parent"] = [parent_category or category.code, parent_code] if parent_code else None
                items.append(item)
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError("فایل Excel معتبر نیست یا از محدودیت اندازه/تعداد ردیف بیشتر است.") from exc
        config = {"format": "crm-engine", "version": 1, "values": items}
        result = import_config(workspace, request.user, config)
        definition_audit(request, workspace, "definition.lookup_excel_imported", category, after=result)
        return Response(result)


class EngineLookupOrderEndpoint(EngineAPIView):
    @transaction.atomic
    def post(self, request, slug, identifier):
        workspace = self.workspace(slug)
        self.require(request, workspace, "Lookup.Edit")
        Workspace.objects.select_for_update().get(id=workspace.id)
        category = get_object_or_404(LookupCategory.objects.select_for_update(), workspace=workspace, id=identifier)
        ids = request.data.get("ids")
        if not isinstance(ids, list) or len(ids) > 10000 or len(ids) != len(set(map(str, ids))):
            raise ValidationError("ترتیب باید شامل شناسه‌های یکتا باشد.")
        rows = list(category.values.all())
        if {str(row.id) for row in rows} != set(ids):
            raise ValidationError("ترتیب باید تمام مقادیر همین لیست را شامل شود.")
        old = [str(row.id) for row in rows]
        positions = {identifier: index for index, identifier in enumerate(ids)}
        for row in rows:
            row.position = positions[str(row.id)]
        LookupValue.objects.bulk_update(rows, ["position"])
        definition_audit(request, workspace, "definition.lookup_reordered", category, before=old, after=ids)
        return Response({"ids": ids})


class EngineFileEndpoint(EngineAPIView):
    def post(self, request, slug, code, field_code):
        entity = self.entity(slug, code)
        if not (
            permission(entity, request.user, "create", request) or permission(entity, request.user, "edit", request)
        ):
            raise PermissionDenied()
        _, fields = schema_fields(
            entity, EntityType.objects.filter(parent=entity, code=request.data.get("subtype")).first()
        )
        field = fields.get(field_code)
        if (
            not field
            or field["kind"] != "file"
            or not can_field(field, role_ids(request.user, entity.workspace), edit=True)
            or not field["is_active"]
            or field["config"].get("hidden")
            or field["config"].get("readonly")
        ):
            raise PermissionDenied()
        upload = request.FILES.get("file")
        config = field["config"]
        if not upload or upload.size > min(config.get("max_file_size", 5 * 1024 * 1024), settings.FILE_SIZE_LIMIT):
            raise ValidationError("فایل انتخاب نشده یا حجم آن بیش از حد مجاز است.")
        extension = PurePath(upload.name).suffix.lower().lstrip(".")
        if config.get("extensions") and extension not in config["extensions"]:
            raise ValidationError("نوع فایل مجاز نیست.")
        asset = FileAsset.objects.create(
            workspace=entity.workspace,
            user=request.user,
            asset=upload,
            entity_type="CRM_FIELD",
            entity_identifier="",
            size=upload.size,
            is_uploaded=True,
            attributes={
                "entity": entity.code,
                "field": field_code,
                "name": PurePath(upload.name).name,
                "content_type": upload.content_type,
            },
        )
        return Response({"id": str(asset.id), "title": asset.attributes["name"]}, status=201)

    def get(self, request, slug, code, field_code, identifier):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "view", request)
        asset = get_object_or_404(
            FileAsset,
            workspace=entity.workspace,
            id=identifier,
            entity_type__in=["CRM_FIELD", "CRM_FIELD_DEFAULT"],
            is_uploaded=True,
            attributes__field=field_code,
        )
        allowed = False
        record_id = request.query_params.get("record") or asset.entity_identifier
        if record_id:
            row, native = get_record(entity, request.user, record_id, request)
            payload = record_payload(entity, request.user, row, native, request)
            allowed = str(asset.id) in (payload["values"].get(field_code) or [])
        elif asset.entity_type == "CRM_FIELD_DEFAULT":
            subtype = EntityType.objects.filter(parent=entity, code=request.query_params.get("subtype")).first()
            field = next(
                (
                    item
                    for item in present_schema(entity, request.user, subtype)["fields"]
                    if item["code"] == field_code
                ),
                None,
            )
            allowed = bool(
                field
                and isinstance(field["config"].get("default"), list)
                and str(asset.id) in field["config"]["default"]
            )
        else:
            allowed = asset.user_id == request.user.id and asset.attributes.get("field") == field_code
        if not allowed:
            raise PermissionDenied()
        return FileResponse(
            asset.asset.open("rb"),
            as_attachment=True,
            filename=asset.attributes.get("name", "file"),
            content_type="application/octet-stream",
        )


class EngineDefaultFileEndpoint(EngineAPIView):
    def field(self, request, slug, identifier, edit=False):
        workspace = self.workspace(slug)
        if not (
            has_permission(request.user, workspace, "FormLayout.Edit", request=request)
            or has_permission(request.user, workspace, "Field.Edit" if edit else "Field.View", request=request)
        ):
            raise PermissionDenied()
        return get_object_or_404(FieldDefinition, workspace=workspace, id=identifier, kind="file", is_active=True)

    def post(self, request, slug, identifier):
        from plane.app.services.entity_files import default_asset

        field = self.field(request, slug, identifier, edit=True)
        upload = request.FILES.get("file")
        if not upload or upload.size > min(
            field.config.get("max_file_size", 5 * 1024 * 1024), settings.FILE_SIZE_LIMIT
        ):
            raise ValidationError("یک فایل در محدوده حجم مجاز انتخاب کنید.")
        asset = default_asset(
            field.workspace, request.user, field.code, field.config, upload.read(), upload.name, upload.content_type
        )
        return Response({"id": str(asset.id), "title": asset.attributes["name"]}, status=201)

    def get(self, request, slug, identifier, asset_id):
        field = self.field(request, slug, identifier)
        asset = get_object_or_404(
            FileAsset,
            workspace=field.workspace,
            id=asset_id,
            entity_type="CRM_FIELD_DEFAULT",
            attributes__field=field.code,
        )
        return FileResponse(
            asset.asset.open("rb"),
            as_attachment=True,
            filename=asset.attributes["name"],
            content_type="application/octet-stream",
        )


class EngineEventsEndpoint(EngineAPIView):
    def get(self, request, slug, code=None, identifier=None):
        workspace = self.workspace(slug)
        rows = EntityEvent.objects.filter(workspace=workspace).select_related("created_by")
        permitted = None
        if code and identifier:
            entity = self.entity(slug, code)
            row, native = get_record(entity, request.user, identifier, request)
            permitted = set(record_payload(entity, request.user, row, native, request)["values"])
            rows = rows.filter(entity=entity, record_id=identifier)
        else:
            self.require(request, workspace, "EntityType.View")
            rows = rows.filter(record_id__isnull=True)
            # Definition auditing must not become a second route to private views/files.
            private_view = Q(metadata__definition_type="issueview") & (
                Q(before__visibility="personal") | Q(after__visibility="personal")
            )
            rows = rows.exclude(private_view & ~Q(created_by=request.user))
            rows = rows.exclude(Q(metadata__definition_type="entityviewexport") & ~Q(created_by=request.user))
            if not has_permission(request.user, workspace, "View.Manage.All", request=request):
                rows = rows.exclude(Q(metadata__definition_type="issueview") & ~Q(created_by=request.user))
        offset, size = page(request)
        result = []
        view_entities = {}
        for event in rows[offset * size : (offset + 1) * size]:
            if permitted is not None and event.field and event.field not in permitted:
                continue
            before, after = event.before, event.after
            if permitted is None and event.event_type.startswith("view."):
                from plane.app.services.entity_views import redact_definition, presentation
                from plane.db.models import IssueView

                key = (event.metadata.get("definition_type"), event.metadata.get("definition_id"))
                if key not in view_entities:
                    if key[0] == "issueview":
                        view = (
                            IssueView.all_objects.filter(workspace=workspace, id=key[1])
                            .select_related("engine_entity__workspace")
                            .first()
                        )
                        view_entities[key] = view.engine_entity if view else None
                    elif key[0] == "entitytype":
                        view_entities[key] = (
                            EntityType.objects.filter(workspace=workspace, id=key[1])
                            .select_related("workspace")
                            .first()
                        )
                    else:
                        view_entities[key] = None
                target = view_entities[key]

                def redact_view_event(value):
                    if not isinstance(value, dict):
                        return value
                    if key[0] == "entitytype" and target:
                        try:
                            return presentation(target, request.user, request, settings=value)
                        except PermissionDenied:
                            return None
                    if "definition" in value:
                        if not target:
                            return None
                        try:
                            return {
                                **value,
                                "definition": redact_definition(target, request.user, value["definition"], request),
                            }
                        except (PermissionDenied, Http404):
                            return None
                    return value

                before, after = redact_view_event(before), redact_view_event(after)
            if permitted is not None and not event.field:
                before = (
                    {k: v for k, v in (before or {}).items() if k in permitted} if isinstance(before, dict) else None
                )
                after = {k: v for k, v in (after or {}).items() if k in permitted} if isinstance(after, dict) else None
            result.append(
                {
                    "id": str(event.id),
                    "type": event.event_type,
                    "field": event.field,
                    "before": before,
                    "after": after,
                    "actor": event.created_by.display_name if event.created_by else "سیستم",
                    "at": event.created_at,
                    "metadata": event.metadata,
                }
            )
        return Response({"results": result, "count": rows.count()})
