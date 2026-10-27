"""Bounded Excel record imports through the same authenticated record writer."""

from datetime import date, datetime
from io import BytesIO
import zipfile
from itertools import islice

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.exceptions import APIException, ValidationError
from rest_framework.response import Response

from plane.app.services.entity_records import require_action, save_record
from plane.app.services.entity_schema import schema_fields, can_field, role_ids
from .entity import EngineAPIView


class EngineRecordImportEndpoint(EngineAPIView):
    def get(self, request, slug, code):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "create", request)
        from openpyxl import Workbook

        _, fields = schema_fields(entity)
        roles = role_ids(request.user, entity.workspace)
        columns = [
            key
            for key, field in fields.items()
            if field["is_active"]
            and not field["hidden_from_form"]
            and not field["config"].get("readonly")
            and field["kind"] != "computed"
            and can_field(field, roles, edit=True)
        ]
        book = Workbook()
        book.active.append(columns)
        output = BytesIO()
        book.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{entity.code}-template.xlsx"'
        )
        return response

    def post(self, request, slug, code):
        entity = self.entity(slug, code)
        require_action(entity, request.user, "create", request)
        upload = request.FILES.get("file")
        if not upload or upload.size > 5 * 1024 * 1024:
            raise ValidationError("فایل Excel با حجم حداکثر پنج مگابایت انتخاب کنید.")
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(upload) as archive:
                if (
                    sum(info.file_size for info in archive.infolist())
                    > 50 * 1024 * 1024
                ):
                    raise ValueError()
            upload.seek(0)
            book = load_workbook(upload, read_only=True, data_only=False)
            rows = iter(book.active.iter_rows(values_only=True))
            headers = list(next(rows))
        except (
            ValueError,
            zipfile.BadZipFile,
            StopIteration,
            KeyError,
            OSError,
        ) as exc:
            raise ValidationError(
                "فایل Excel معتبر نیست یا از محدودیت اندازه بیشتر است."
            ) from exc
        if (
            not headers
            or len(headers) > 300
            or any(not isinstance(header, str) for header in headers)
            or len(set(headers)) != len(headers)
        ):
            raise ValidationError(
                "ردیف نخست باید کد ستون‌های غیرتکراری را داشته باشد؛ از الگوی ورود داده استفاده کنید."
            )
        subtype_code = request.data.get("subtype")
        subtype = (
            entity.subtypes.filter(code=subtype_code, is_active=True).first()
            if subtype_code
            else None
        )
        if subtype_code and not subtype:
            raise ValidationError("زیرنوع معتبر نیست.")
        _, fields = schema_fields(entity, subtype)
        unknown = set(headers) - set(fields)
        if unknown:
            raise ValidationError(
                "ستون ناشناخته در فایل وجود دارد؛ الگوی همین موجودیت را دریافت کنید."
            )
        results = []
        pending = list(islice(rows, 10001))
        if len(pending) > 10000:
            book.close()
            raise ValidationError(
                "هر فایل حداکثر ده هزار ردیف دارد؛ فایل به چند بخش تقسیم شود."
            )
        try:
            for index, row in enumerate(pending, 2):
                if not any(value is not None for value in row):
                    continue
                values, warnings = {}, []
                for key, value in zip(headers, row):
                    field = fields[key]
                    if field["hidden_from_form"] or field["config"].get("auto_number"):
                        warnings.append(
                            f"ستون {key} نادیده گرفته شد؛ مقدار جدید نمی‌پذیرد."
                        )
                        continue
                    if isinstance(value, str) and value.startswith("="):
                        values = None
                        break
                    if isinstance(value, datetime):
                        value = (
                            value.date().isoformat()
                            if field["kind"] == "date"
                            else (
                                timezone.make_aware(value).isoformat()
                                if timezone.is_naive(value)
                                else value.isoformat()
                            )
                        )
                    elif isinstance(value, date):
                        value = value.isoformat()
                    values[key] = value
                if values is None:
                    results.append(
                        {
                            "row": index,
                            "ok": False,
                            "error": "فرمول در فایل ورودی پذیرفته نیست.",
                            "warnings": warnings,
                        }
                    )
                    continue
                try:
                    record = save_record(
                        entity,
                        request.user,
                        {
                            "values": values,
                            **({"subtype": subtype_code} if subtype_code else {}),
                        },
                        request,
                    )
                    results.append(
                        {
                            "row": index,
                            "ok": True,
                            "id": record["id"],
                            "number": record["number"],
                            "warnings": [*warnings, *record.get("warnings", [])],
                        }
                    )
                except APIException as exc:
                    results.append(
                        {
                            "row": index,
                            "ok": False,
                            "error": exc.detail,
                            "warnings": warnings,
                        }
                    )
        finally:
            book.close()
        return Response(
            {
                "results": results,
                "success_count": sum(result["ok"] for result in results),
                "rejected_count": sum(not result["ok"] for result in results),
            }
        )
