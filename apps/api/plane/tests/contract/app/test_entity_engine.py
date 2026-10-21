"""Behavioral contracts for Ticket 5A: metadata, published forms and real records."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from django.core.files.uploadedfile import SimpleUploadedFile

from plane.db.models import (
    EntityType,
    FieldDefinition,
    FormLayout,
    FormVersion,
    EntityRecord,
    EntityEvent,
    LookupCategory,
    LookupValue,
    OrganizationRole,
    UserOrganizationRole,
    Workspace,
    WorkspaceMember,
    User,
)
from plane.app.services.entity_schema import initialize_engine
from plane.app.services.entity_config import export_config, import_config
from plane.app.services.permission_registry import sync_permission_catalog


pytestmark = [pytest.mark.django_db, pytest.mark.contract]


@pytest.fixture
def engine(workspace, create_user, session_client):
    initialize_engine(workspace)
    sync_permission_catalog(workspace)
    return workspace, create_user, session_client


def endpoint(workspace, suffix):
    return f"/api/workspaces/{workspace.slug}/engine/{suffix}"


def make_entity(engine, fields=None, code="sample", extra_bindings=None):
    workspace, actor, client = engine
    entity = EntityType.objects.create(
        workspace=workspace, code=code, title="آزمایش", numbering={"prefix": "S", "padding": 4}
    )
    form_fields = [{"field": "title"}]
    for definition in fields or []:
        raw = dict(definition)
        binding = raw.pop("binding", {})
        field = FieldDefinition.objects.create(workspace=workspace, **raw)
        form_fields.append({"field": field.code, **binding})
    form_fields += extra_bindings or []
    layout = {"tabs": [], "sections": [{"id": "main", "title": "اصلی", "columns": 2, "bindings": form_fields}]}
    form = FormLayout.objects.create(entity=entity, draft=layout)
    version = FormVersion.objects.create(layout=form, number=1, schema=layout)
    form.published = version
    form.save()
    return entity, form


def create_record(engine, entity, values, expected=201, **payload):
    workspace, _, client = engine
    response = client.post(
        endpoint(workspace, f"entities/{entity.code}/records/"), {"values": values, **payload}, format="json"
    )
    assert response.status_code == expected, response.data
    return response.data


@pytest.mark.django_db(transaction=True)
def test_concurrent_numbering_and_unique_values_are_serialized(engine):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    from django.db import connections
    from rest_framework.test import APIClient

    workspace, actor, _ = engine
    entity, _ = make_entity(
        engine, [{"code": "unique_key", "title": "کلید یکتا", "kind": "short_text", "config": {"unique": True}}]
    )
    url = endpoint(workspace, f"entities/{entity.code}/records/")

    def submit(key, start):
        try:
            client = APIClient()
            client.force_authenticate(user=User.objects.get(id=actor.id))
            start.wait(timeout=30)
            response = client.post(url, {"values": {"title": key, "unique_key": key}}, format="json")
            return response.status_code, response.data
        finally:
            connections.close_all()

    with ThreadPoolExecutor(max_workers=4) as pool:
        start = Barrier(4)
        attempts = [pool.submit(submit, f"key-{index}", start) for index in range(4)]
        responses = [attempt.result(timeout=90) for attempt in attempts]
        assert [status for status, _ in responses] == [201] * 4, responses
        assert {data["number"] for _, data in responses} == {"S-0001", "S-0002", "S-0003", "S-0004"}
        start = Barrier(2)
        attempts = [pool.submit(submit, "same-key", start) for _ in range(2)]
        responses = [attempt.result(timeout=90) for attempt in attempts]
        assert sorted(status for status, _ in responses) == [201, 400], responses
    assert EntityRecord.objects.filter(entity=entity, data__unique_key="same-key").count() == 1
    assert EntityRecord.objects.filter(entity=entity).count() == 5


def test_rolled_back_records_do_not_leave_success_events_or_audit_receipts(engine):
    from django.db import transaction

    from plane.app.services.entity_records import save_record
    from plane.db.models import APIActivityLog

    workspace, actor, _ = engine
    entity, _ = make_entity(engine)
    receipts = APIActivityLog.objects.filter(token_identifier=f"crm:{workspace.id}:{actor.id}")
    initial_receipts = receipts.count()
    with pytest.raises(RuntimeError, match="rollback probe"):
        with transaction.atomic():
            save_record(entity, actor, {"values": {"title": "rolled back"}})
            assert EntityEvent.objects.filter(entity=entity, event_type="record.created").exists()
            assert receipts.count() > initial_receipts
            raise RuntimeError("rollback probe")
    assert not EntityRecord.objects.filter(entity=entity).exists()
    assert not EntityEvent.objects.filter(entity=entity).exists()
    assert receipts.count() == initial_receipts


@pytest.mark.parametrize(
    "kind,value,expected",
    [
        ("short_text", "متن", "متن"),
        ("long_text", "متن بلند", "متن بلند"),
        ("rich_text", "<p>سلام</p>", "<p>سلام</p>"),
        ("integer", 0, 0),
        ("decimal", "12.345", "12.345"),
        ("money", "123456789.12", "123456789.12"),
        ("percent", 25, 25.0),
        ("date", "2026-10-19", "2026-10-19"),
        ("datetime", "2026-10-19T10:00:00+03:30", "2026-10-19T10:00:00+03:30"),
        ("date_range", {"start": "2026-10-01", "end": "2026-10-20"}, {"start": "2026-10-01", "end": "2026-10-20"}),
        ("checkbox", False, False),
        ("boolean", True, True),
        ("tags", ["الف", "ب"], ["الف", "ب"]),
        ("url", "https://example.com/path", "https://example.com/path"),
        ("email", "user@example.com", "user@example.com"),
        ("phone", "+98 912 1234567", "+98 912 1234567"),
        ("rating", 4, 4),
    ],
)
def test_typed_values_round_trip_and_preserve_false_zero(engine, kind, value, expected):
    entity, _ = make_entity(engine, [{"code": "test_value", "title": "آزمایش", "kind": kind}])
    data = create_record(engine, entity, {"title": "اول", "test_value": value})
    assert data["values"]["test_value"] == expected
    response = engine[2].get(endpoint(engine[0], f"entities/{entity.code}/records/{data['id']}/"))
    assert response.status_code == 200
    assert response.data["values"]["test_value"] == expected


def test_configuration_only_crud_defaults_events_and_partial_update(engine):
    workspace, actor, client = engine
    entity, _ = make_entity(
        engine,
        [
            {"code": "owner", "title": "مالک", "kind": "user", "config": {"default_dynamic": "current_user"}},
            {"code": "today", "title": "امروز", "kind": "date", "config": {"default_dynamic": "today"}},
            {"code": "amount", "title": "مبلغ", "kind": "integer", "config": {"min": 0, "default": 3}},
            {
                "code": "total",
                "title": "جمع",
                "kind": "computed",
                "config": {"formula": {"op": "multiply", "args": [{"field": "amount"}, {"value": 2}]}},
            },
        ],
    )
    created = create_record(engine, entity, {"title": "اول"})
    assert created["number"] == "S-0001"
    assert created["values"]["owner"] == str(actor.id)
    assert created["values"]["amount"] == 3
    assert created["values"]["total"] == 6
    response = client.patch(
        endpoint(workspace, f"entities/sample/records/{created['id']}/"),
        {"revision": created["revision"], "values": {"amount": 5}},
        format="json",
    )
    assert response.status_code == 200, response.data
    assert response.data["values"]["title"] == "اول"
    assert response.data["values"]["total"] == 10
    assert EntityEvent.objects.filter(entity=entity, event_type="record.field_changed", field="amount").exists()
    assert (
        client.patch(
            endpoint(workspace, f"entities/sample/records/{created['id']}/"),
            {"revision": 1, "values": {"title": "قدیمی"}},
            format="json",
        ).status_code
        == 409
    )


def test_conditional_required_visibility_and_validation_cannot_be_bypassed(engine):
    entity, _ = make_entity(
        engine,
        [
            {"code": "urgent", "title": "فوری", "kind": "boolean"},
            {
                "code": "urgent_reason",
                "title": "علت",
                "kind": "short_text",
                "config": {"min_length": 3},
                "binding": {
                    "required_if": {"field": "urgent", "op": "eq", "value": True},
                    "visible_if": {"field": "urgent", "op": "eq", "value": True},
                },
            },
        ],
    )
    create_record(engine, entity, {"title": "اول", "urgent": True}, 400)
    create_record(engine, entity, {"title": "اول", "urgent": False, "urgent_reason": "غیرمجاز"}, 403)
    create_record(engine, entity, {"title": "اول", "urgent": True, "urgent_reason": "کافی"})
    create_record(engine, entity, {"title": "اول", "urgent": True, "urgent_reason": "کم"}, 400)


def test_readonly_frozen_values_limits_and_regex_cannot_be_bypassed_by_patch(engine):
    workspace, _, client = engine
    entity, _ = make_entity(
        engine,
        [
            {"code": "frozen", "title": "ثابت", "kind": "short_text", "config": {"editable_after_create": False}},
            {
                "code": "system",
                "title": "سیستمی",
                "kind": "short_text",
                "config": {"readonly": True, "default": "auto"},
            },
            {"code": "amount", "title": "عدد", "kind": "integer", "config": {"min": 0, "max": 10}},
            {
                "code": "pattern",
                "title": "الگو",
                "kind": "short_text",
                "config": {
                    "min_length": 2,
                    "max_length": 3,
                    "regex": "[A-Z]{2,3}",
                    "regex_message": "دو یا سه حرف بزرگ",
                },
            },
        ],
    )
    data = create_record(engine, entity, {"title": "valid", "frozen": "original", "amount": 0, "pattern": "AB"})
    assert data["values"]["system"] == "auto"
    url = endpoint(workspace, f"entities/{entity.code}/records/{data['id']}/")
    for values, status in [
        ({"frozen": "changed"}, 400),
        ({"system": "changed"}, 403),
        ({"amount": -1}, 400),
        ({"amount": 11}, 400),
        ({"pattern": "A"}, 400),
        ({"pattern": "ABCD"}, 400),
        ({"pattern": "ab"}, 400),
    ]:
        response = client.patch(url, {"revision": data["revision"], "values": values}, format="json")
        assert response.status_code == status, response.data
        if values == {"pattern": "ab"}:
            assert "دو یا سه حرف بزرگ" in str(response.data)
    updated = client.patch(url, {"revision": data["revision"], "values": {"title": "changed"}}, format="json")
    assert updated.status_code == 200, updated.data
    assert updated.data["values"] == {**data["values"], "title": "changed"}
    field = FieldDefinition.objects.get(workspace=workspace, code="amount")
    assert (
        client.patch(
            endpoint(workspace, f"definitions/fields/{field.id}/"), {"config": {"min": -10}}, format="json"
        ).status_code
        == 400
    )


def test_used_field_immutable_and_unbound_data_survives_publish(engine):
    workspace, _, client = engine
    entity, form = make_entity(engine, [{"code": "keep", "title": "قدیمی", "kind": "short_text"}])
    record = create_record(engine, entity, {"title": "اول", "keep": "حفظ شود"})
    field = FieldDefinition.objects.get(workspace=workspace, code="keep")
    url = endpoint(workspace, f"definitions/fields/{field.id}/")
    assert client.patch(url, {"kind": "integer"}, format="json").status_code == 400
    assert client.delete(url).status_code == 400
    assert client.patch(url, {"title": "نام تازه"}, format="json").status_code == 200
    draft = {"tabs": [], "sections": [{"id": "main", "columns": 1, "bindings": [{"field": "title"}]}]}
    saved = client.put(endpoint(workspace, "forms/sample/"), {"revision": form.revision, "draft": draft}, format="json")
    assert saved.status_code == 200
    assert (
        client.post(
            endpoint(workspace, "forms/sample/"), {"revision": saved.data["revision"]}, format="json"
        ).status_code
        == 200
    )
    response = client.patch(
        endpoint(workspace, f"entities/sample/records/{record['id']}/"),
        {"revision": 1, "values": {"title": "جدید"}},
        format="json",
    )
    assert response.status_code == 200, response.data
    assert EntityRecord.objects.get(id=record["id"]).data["keep"] == "حفظ شود"


def test_lookup_cascading_historical_deactivation_and_delete_policy(engine):
    workspace, _, client = engine
    category = LookupCategory.objects.create(workspace=workspace, code="cities", title="شهرها")
    parent_category = LookupCategory.objects.create(workspace=workspace, code="provinces", title="استان‌ها")
    parent = LookupValue.objects.create(category=parent_category, code="tehran", title="تهران")
    value = LookupValue.objects.create(category=category, code="city", title="شهر", parent=parent)
    entity, _ = make_entity(
        engine,
        [
            {"code": "province", "title": "استان", "kind": "select", "lookup": parent_category},
            {
                "code": "city",
                "title": "شهر",
                "kind": "select",
                "lookup": category,
                "config": {"depends_on": "province"},
            },
        ],
    )
    create_record(engine, entity, {"title": "اول", "city": str(value.id)}, 400)
    old = create_record(engine, entity, {"title": "اول", "province": str(parent.id), "city": str(value.id)})
    another_parent = LookupValue.objects.create(category=parent_category, code="other", title="استان دیگر")
    changed_parent = client.patch(
        endpoint(workspace, f"entities/sample/records/{old['id']}/"),
        {"revision": old["revision"], "values": {"province": str(another_parent.id)}},
        format="json",
    )
    assert changed_parent.status_code == 400, changed_parent.data
    url = endpoint(workspace, f"definitions/values/{value.id}/")
    assert client.delete(url).status_code == 400
    assert client.patch(url, {"is_active": False}, format="json").status_code == 200
    create_record(engine, entity, {"title": "دوم", "province": str(parent.id), "city": str(value.id)}, 400)
    assert client.get(endpoint(workspace, f"entities/sample/records/{old['id']}/")).data["values"]["city"] == str(
        value.id
    )
    assert client.patch(url, {"is_active": True}, format="json").status_code == 200
    unused = LookupValue.objects.create(category=category, code="unused", title="بدون استفاده")
    assert client.delete(endpoint(workspace, f"definitions/values/{unused.id}/")).status_code == 204


def test_binding_role_acl_enforced_on_schema_read_and_write(engine):
    workspace, owner, client = engine
    secret_role = OrganizationRole.objects.create(workspace=workspace, name="مالی")
    entity, _ = make_entity(
        engine,
        [
            {
                "code": "secret",
                "title": "محرمانه",
                "kind": "money",
                "binding": {"view_roles": [str(secret_role.id)], "edit_roles": [str(secret_role.id)]},
            }
        ],
    )
    # Built-in administrators do not implicitly override explicit field restrictions.
    create_record(engine, entity, {"title": "اول", "secret": "100"}, 403)
    schema = client.get(endpoint(workspace, "entities/sample/schema/"))
    assert "secret" not in {f["code"] for f in schema.data["fields"]}
    UserOrganizationRole.objects.create(workspace=workspace, user=owner, role=secret_role)
    created = create_record(engine, entity, {"title": "اول", "secret": "100"})
    UserOrganizationRole.objects.filter(workspace=workspace, user=owner, role=secret_role).delete()
    response = client.get(endpoint(workspace, f"entities/sample/records/{created['id']}/"))
    assert "secret" not in response.data["values"]
    history = client.get(endpoint(workspace, f"entities/sample/records/{created['id']}/events/"))
    assert all(event["field"] != "secret" for event in history.data["results"])
    assert all("secret" not in (event["after"] or {}) for event in history.data["results"] if not event["field"])


def test_draft_publish_rollback_and_concurrent_edit(engine):
    workspace, _, client = engine
    entity, form = make_entity(engine)
    url = endpoint(workspace, "forms/sample/")
    draft = {
        "tabs": [{"id": "tab", "title": "جدید"}],
        "sections": [
            {
                "id": "s",
                "title": "بخش",
                "tab": "tab",
                "columns": 3,
                "bindings": [{"field": "title", "title": "عنوان تازه"}],
            }
        ],
    }
    saved = client.put(url, {"revision": 1, "draft": draft}, format="json")
    assert saved.status_code == 200, saved.data
    assert client.get(endpoint(workspace, "entities/sample/schema/")).data["fields"][0]["title"] == "عنوان"
    assert client.put(url, {"revision": 1, "draft": draft}, format="json").status_code == 409
    published = client.post(url, {"revision": saved.data["revision"], "change_note": "انتشار"}, format="json")
    assert published.status_code == 200, published.data
    assert client.get(endpoint(workspace, "entities/sample/schema/")).data["fields"][0]["title"] == "عنوان تازه"
    rollback = client.post(url, {"revision": published.data["revision"], "rollback_version": 1}, format="json")
    assert rollback.status_code == 200
    assert rollback.data["version"] == 3
    assert FormVersion.objects.filter(layout=form).count() == 3


def test_relation_server_search_reverse_and_confirmation(engine):
    workspace, _, client = engine
    target, _ = make_entity(engine, code="target")
    target_row = create_record(engine, target, {"title": "مقصد قابل جستجو"})
    source, _ = make_entity(
        engine, [{"code": "target_link", "title": "مقصد", "kind": "relation", "relation_entity": target}], code="source"
    )
    row = create_record(engine, source, {"title": "مبدا", "target_link": target_row["id"]})
    options = client.get(endpoint(workspace, "entities/source/options/target_link/?search=مقصد"))
    assert options.status_code == 200
    assert options.data["results"][0]["id"] == target_row["id"]
    inverse = client.get(endpoint(workspace, f"entities/target/records/{target_row['id']}/relations/"))
    assert inverse.data["results"][0]["id"] == row["id"]
    assert client.delete(endpoint(workspace, f"entities/target/records/{target_row['id']}/")).status_code == 409


def test_excel_values_round_trip_and_import_is_atomic(engine):
    workspace, _, client = engine
    category = LookupCategory.objects.create(workspace=workspace, code="excel", title="اکسل")
    url = endpoint(workspace, f"lookups/{category.id}/excel/")
    book = Workbook()
    book.active.append(["code", "title", "color", "position"])
    book.active.append(["first", "اول", "#123456", 1])
    stream = BytesIO()
    book.save(stream)
    response = client.post(url, {"file": SimpleUploadedFile("values.xlsx", stream.getvalue())}, format="multipart")
    assert response.status_code == 200, response.data
    response = client.get(url)
    exported = load_workbook(BytesIO(response.content))
    assert exported.active["B2"].value == "اول"
    assert LookupValue.objects.get(category=category).color == "#123456"


def test_templates_and_portable_configuration(engine):
    workspace, actor, client = engine
    for template in ("support", "sales", "hr"):
        response = client.post(endpoint(workspace, "config/"), {"template": template}, format="json")
        assert response.status_code == 200, response.data
        assert client.post(endpoint(workspace, "config/"), {"template": template}, format="json").status_code == 200
    exported = export_config(workspace)
    other = Workspace.objects.create(name="دیگر", slug="other", owner=actor)
    WorkspaceMember.objects.create(workspace=other, member=actor, role=20)
    initialize_engine(other)
    result = import_config(other, actor, exported)
    assert result["forms"] >= 7
    assert FieldDefinition.objects.get(workspace=other, code="sales_amount").kind == "money"
    assert EntityType.objects.get(workspace=other, code="template-hr").parent.code == "leave"


def test_cross_workspace_and_definition_permissions(engine):
    workspace, actor, client = engine
    stranger = User.objects.create_user(
        email="stranger-entity@example.com", username="stranger", password="SafePassword!2026"
    )
    WorkspaceMember.objects.create(workspace=workspace, member=stranger, role=15)
    entity, _ = make_entity(engine)
    row = create_record(engine, entity, {"title": "محرمانه"})
    client.force_authenticate(user=stranger)
    assert (
        client.post(
            endpoint(workspace, "definitions/entities/"), {"code": "forbidden", "title": "ممنوع"}, format="json"
        ).status_code
        == 403
    )
    assert client.get(endpoint(workspace, f"entities/sample/records/{row['id']}/")).status_code == 404
    other = Workspace.objects.create(name="بیگانه", slug="outsider", owner=actor)
    assert client.get(endpoint(other, "catalog/")).status_code == 403


def test_formula_cycle_unknown_field_and_invalid_regex_rejected_before_publish(engine):
    workspace, _, client = engine
    entity, form = make_entity(engine)
    invalid = {"sections": [{"id": "x", "columns": 2, "bindings": [{"field": "missing"}]}]}
    assert (
        client.put(endpoint(workspace, "forms/sample/"), {"revision": 1, "draft": invalid}, format="json").status_code
        == 400
    )
    response = client.post(
        endpoint(workspace, "definitions/fields/"),
        {"code": "bad_regex", "title": "الگو", "kind": "short_text", "config": {"regex": "(a+)+$"}},
        format="json",
    )
    assert response.status_code == 400


def test_native_project_creation_uses_existing_model(engine):
    from plane.db.models import Project, ProjectMember, State

    workspace, actor, client = engine
    entity = EntityType.objects.get(workspace=workspace, code="project")
    record = create_record(engine, entity, {"title": "پروژه موتور", "description": "متن"})
    native = Project.objects.get(id=record["id"])
    assert native.name == "پروژه موتور"
    assert ProjectMember.objects.filter(project=native, member=actor).exists()
    assert State.objects.filter(project=native).exists()
    assert EntityRecord.objects.get(entity=entity, source_id=native.id).data == {}


def test_native_api_cannot_bypass_required_custom_fields_and_preserves_them(engine):
    from plane.db.models import Project

    workspace, actor, client = engine
    entity = EntityType.objects.get(workspace=workspace, code="project")
    field = FieldDefinition.objects.create(
        workspace=workspace, code="project_reference", title="مرجع", kind="short_text", config={"required": True}
    )
    form = FormLayout.objects.get(entity=entity)
    form.draft["sections"][0]["bindings"].append({"field": field.code})
    version = FormVersion.objects.create(layout=form, number=2, schema=form.draft)
    form.published = version
    form.save()
    native_url = f"/api/workspaces/{workspace.slug}/projects/"
    failed = client.post(native_url, {"name": "پروژه بدون مرجع"}, format="json")
    assert failed.status_code == 400, failed.data
    assert not Project.objects.filter(workspace=workspace, name="پروژه بدون مرجع").exists()
    created = client.post(
        native_url, {"name": "پروژه با مرجع", "custom_fields": {"project_reference": "REF-1"}}, format="json"
    )
    assert created.status_code == 201, created.data
    sidecar = EntityRecord.objects.get(entity=entity, source_id=created.data["id"])
    assert sidecar.data["project_reference"] == "REF-1"
    assert EntityEvent.objects.filter(entity=entity, record_id=created.data["id"], event_type="record.created").exists()


def test_native_field_acl_applies_to_old_api_read_write_and_search(engine):
    workspace, actor, client = engine
    entity = EntityType.objects.get(workspace=workspace, code="project")
    record = create_record(engine, entity, {"title": "عنوان محرمانه"})
    form = FormLayout.objects.get(entity=entity)
    secret = OrganizationRole.objects.create(workspace=workspace, name="نقش محرمانه")
    form.draft["sections"][0]["bindings"][0]["view_roles"] = [str(secret.id)]
    form.draft["sections"][0]["bindings"][0]["edit_roles"] = [str(secret.id)]
    form.published = FormVersion.objects.create(layout=form, number=2, schema=form.draft)
    form.save()
    response = client.get(f"/api/workspaces/{workspace.slug}/projects/{record['id']}/")
    assert response.status_code == 200, response.data
    assert "name" not in response.data
    assert (
        client.patch(
            f"/api/workspaces/{workspace.slug}/projects/{record['id']}/", {"name": "دور زدن"}, format="json"
        ).status_code
        == 403
    )
    assert client.get(f"/api/workspaces/{workspace.slug}/projects/?search=محرمانه").status_code == 403


def test_lookup_and_binding_defaults_transfer_by_code_not_database_uuid(engine):
    workspace, actor, _ = engine
    category = LookupCategory.objects.create(workspace=workspace, code="portable", title="قابل انتقال")
    option = LookupValue.objects.create(category=category, code="stable", title="پایدار")
    entity, form = make_entity(
        engine,
        [
            {
                "code": "portable_choice",
                "title": "انتخاب",
                "kind": "select",
                "lookup": category,
                "config": {"default": str(option.id)},
            }
        ],
    )
    form.draft["sections"][0]["bindings"][1]["default"] = str(option.id)
    form.save()
    other = Workspace.objects.create(name="مقصد", slug="portable-target", owner=actor)
    WorkspaceMember.objects.create(workspace=other, member=actor, role=20)
    import_config(other, actor, export_config(workspace))
    destination = LookupValue.objects.get(category__workspace=other, category__code="portable", code="stable")
    assert destination.id != option.id
    field = FieldDefinition.objects.get(workspace=other, code="portable_choice")
    assert field.config["default"] == str(destination.id)
    assert FormLayout.objects.get(entity__workspace=other, entity__code="sample").draft["sections"][0]["bindings"][1][
        "default"
    ] == str(destination.id)


def test_multi_lookup_relation_user_unit_role_and_file_round_trip(engine, tmp_path, monkeypatch):
    from django.core.files.storage import FileSystemStorage
    from plane.db.models import OrganizationUnit, FileAsset

    workspace, actor, client = engine
    category = LookupCategory.objects.create(workspace=workspace, code="multiple", title="چندانتخابی")
    options = [LookupValue.objects.create(category=category, code=f"o{i}", title=f"گزینه {i}") for i in range(2)]
    role = OrganizationRole.objects.create(workspace=workspace, name="نقش آزمایشی")
    unit = OrganizationUnit.objects.create(workspace=workspace, title="واحد آزمایشی")
    target, _ = make_entity(engine, code="related")
    target_record = create_record(engine, target, {"title": "مقصد"})
    entity, _ = make_entity(
        engine,
        [
            {"code": "choices", "title": "گزینه‌ها", "kind": "multi_select", "lookup": category},
            {
                "code": "related_records",
                "title": "ارتباط‌ها",
                "kind": "relation",
                "relation_entity": target,
                "config": {"multiple": True},
            },
            {"code": "person", "title": "کاربر", "kind": "user"},
            {"code": "department", "title": "واحد", "kind": "unit"},
            {"code": "job", "title": "نقش", "kind": "role"},
            {
                "code": "document",
                "title": "فایل",
                "kind": "file",
                "config": {"extensions": ["txt"], "max_file_size": 1000},
            },
        ],
    )
    monkeypatch.setattr(FileAsset._meta.get_field("asset"), "storage", FileSystemStorage(location=tmp_path))
    uploaded = client.post(
        endpoint(workspace, "entities/sample/files/document/"),
        {"file": SimpleUploadedFile("note.txt", b"example")},
        format="multipart",
    )
    assert uploaded.status_code == 201, uploaded.data
    payload = {
        "title": "همه انتخاب‌ها",
        "choices": [str(v.id) for v in options],
        "related_records": [target_record["id"]],
        "person": str(actor.id),
        "department": str(unit.id),
        "job": str(role.id),
        "document": [uploaded.data["id"]],
    }
    record = create_record(engine, entity, payload)
    assert record["values"] == payload
    download = client.get(endpoint(workspace, f"entities/sample/files/document/{uploaded.data['id']}/"))
    assert download.status_code == 200
    assert b"".join(download.streaming_content) == b"example"
    invalid = client.post(
        endpoint(workspace, "entities/sample/files/document/"),
        {"file": SimpleUploadedFile("bad.exe", b"bad")},
        format="multipart",
    )
    assert invalid.status_code == 400


def test_unused_category_and_entity_can_be_deleted_without_losing_audit(engine):
    workspace, _, client = engine
    category = LookupCategory.objects.create(workspace=workspace, code="disposable", title="بدون استفاده")
    LookupValue.objects.create(category=category, code="unused", title="بدون استفاده")
    assert client.delete(endpoint(workspace, f"definitions/lookups/{category.id}/")).status_code == 204
    entity, form = make_entity(engine)
    assert client.delete(endpoint(workspace, f"definitions/entities/{entity.id}/")).status_code == 204
    assert EntityEvent.objects.filter(event_type="definition.entities.deleted").exists()


def test_native_ticket_task_leave_and_dynamic_business_choices(engine):
    from django.utils import timezone
    from datetime import timedelta
    from plane.db.models import Issue, WorkspaceTask, UserAbsence

    workspace, actor, client = engine
    project = create_record(
        engine, EntityType.objects.get(workspace=workspace, code="project"), {"title": "پروژه واقعی"}
    )
    priority = LookupCategory.objects.get(workspace=workspace, code="priority")
    LookupValue.objects.create(category=priority, code="custom", title="اولویت تازه")
    ticket = create_record(
        engine,
        EntityType.objects.get(workspace=workspace, code="ticket"),
        {"title": "تیکت واقعی", "project": project["id"], "ticket_assignees": [str(actor.id)], "priority": "custom"},
    )
    assert Issue.objects.get(id=ticket["id"]).priority == "custom"
    assert ticket["values"]["ticket_assignees"] == [str(actor.id)]
    statuses = LookupCategory.objects.get(workspace=workspace, code="task-status")
    LookupValue.objects.create(category=statuses, code="custom_status", title="وضعیت تازه")
    task = create_record(
        engine,
        EntityType.objects.get(workspace=workspace, code="task"),
        {"title": "وظیفه واقعی", "assignee": str(actor.id), "task_status": "custom_status"},
    )
    assert WorkspaceTask.objects.get(id=task["id"]).status == "custom_status"
    leave = create_record(
        engine,
        EntityType.objects.get(workspace=workspace, code="leave"),
        {
            "leave_user": str(actor.id),
            "starts_at": (timezone.now() + timedelta(days=2)).isoformat(),
            "ends_at": (timezone.now() + timedelta(days=3)).isoformat(),
            "reason": "آزمایش فرم مرخصی",
        },
    )
    assert UserAbsence.objects.get(id=leave["id"]).user_id == actor.id


def test_ticket_subtype_extends_native_issue_type_and_keeps_existing_types(engine):
    from plane.db.models import IssueType, Issue

    workspace, _, client = engine
    root = EntityType.objects.get(workspace=workspace, code="ticket")
    existing = IssueType.objects.create(
        workspace=workspace, name="نوع قبلی", external_source="integration", external_id="keep-me"
    )
    mirrored = EntityType.objects.get(workspace=workspace, code=f"issue-type-{existing.id}")
    assert mirrored.parent == root
    existing.refresh_from_db()
    assert existing.external_source == "integration" and existing.external_id == "keep-me"
    subtype = client.post(
        endpoint(workspace, "definitions/entities/"),
        {"parent": str(root.id), "code": "complaint", "title": "شکایت"},
        format="json",
    )
    assert subtype.status_code == 201, subtype.data
    native_type = IssueType.objects.get(
        workspace=workspace, external_source="crm-entity-engine", external_id="complaint"
    )
    project = create_record(
        engine, EntityType.objects.get(workspace=workspace, code="project"), {"title": "پروژه زیرنوع"}
    )
    ticket = create_record(engine, root, {"title": "شکایت واقعی", "project": project["id"]}, subtype="complaint")
    assert Issue.objects.get(id=ticket["id"]).type_id == native_type.id
    assert ticket["subtype"] == "complaint"


def test_hidden_field_defaults_and_computed_dependencies_do_not_leak(engine):
    workspace, _, client = engine
    entity, _ = make_entity(
        engine,
        [
            {"code": "secret_number", "title": "پنهان", "kind": "integer", "config": {"hidden": True, "default": 42}},
            {
                "code": "secret_result",
                "title": "محاسبه",
                "kind": "computed",
                "config": {"formula": {"field": "secret_number"}},
            },
        ],
    )
    schema = client.get(endpoint(workspace, "entities/sample/schema/"))
    assert {field["code"] for field in schema.data["fields"]} == {"title"}
    record = create_record(engine, entity, {"title": "رکورد"})
    assert "secret_number" not in record["values"] and "secret_result" not in record["values"]


def test_relation_payload_hides_target_when_target_scope_is_not_granted(engine):
    workspace, _, client = engine
    target, _ = make_entity(engine, code="private-target")
    hidden_target = create_record(engine, target, {"title": "مقصد خصوصی"})
    source, _ = make_entity(
        engine,
        [{"code": "private_link", "title": "ارتباط", "kind": "relation", "relation_entity": target}],
        code="visible-source",
    )
    source_record = create_record(engine, source, {"title": "مبدا", "private_link": hidden_target["id"]})
    viewer = User.objects.create_user(email="relation-limited@example.com", username="relation-limited")
    WorkspaceMember.objects.create(workspace=workspace, member=viewer, role=15)
    EntityRecord.objects.filter(id=source_record["id"]).update(created_by=viewer)
    client.force_authenticate(viewer)
    response = client.get(endpoint(workspace, f"entities/visible-source/records/{source_record['id']}/"))
    assert response.status_code == 200
    assert response.data["values"]["private_link"] is None
    assert EntityRecord.objects.get(id=source_record["id"]).data["private_link"] == hidden_target["id"]


def test_default_attachment_reuse_permissions_and_portable_bytes(engine, tmp_path, monkeypatch):
    from django.core.files.storage import FileSystemStorage
    from plane.db.models import FileAsset

    workspace, owner, client = engine
    monkeypatch.setattr(FileAsset._meta.get_field("asset"), "storage", FileSystemStorage(location=tmp_path))
    entity, _ = make_entity(
        engine,
        [
            {
                "code": "default_document",
                "title": "پیوست پیش‌فرض",
                "kind": "file",
                "config": {"extensions": ["txt"], "max_file_size": 1000},
            }
        ],
    )
    field = FieldDefinition.objects.get(workspace=workspace, code="default_document")
    uploaded = client.post(
        endpoint(workspace, f"definitions/fields/{field.id}/default-file/"),
        {"file": SimpleUploadedFile("guide.txt", b"reusable default")},
        format="multipart",
    )
    assert uploaded.status_code == 201, uploaded.data
    configured = client.patch(
        endpoint(workspace, f"definitions/fields/{field.id}/"),
        {"config": {**field.config, "default": [uploaded.data["id"]]}},
        format="json",
    )
    assert configured.status_code == 200, configured.data
    first = create_record(engine, entity, {"title": "اول"})
    viewer = User.objects.create_user(email="default-file-member@example.com", username="default-file-member")
    WorkspaceMember.objects.create(workspace=workspace, member=viewer, role=15)
    client.force_authenticate(viewer)
    second = create_record(engine, entity, {"title": "دوم"})
    assert first["values"][field.code] == second["values"][field.code] == [uploaded.data["id"]]
    url = endpoint(workspace, f"entities/sample/files/{field.code}/{uploaded.data['id']}/")
    assert client.get(url + f"?record={first['id']}").status_code == 404
    downloaded = client.get(url + f"?record={second['id']}")
    assert downloaded.status_code == 200
    assert b"".join(downloaded.streaming_content) == b"reusable default"
    assert FileAsset.objects.get(id=uploaded.data["id"]).entity_identifier == ""
    client.force_authenticate(owner)
    config = export_config(workspace)
    other = Workspace.objects.create(name="مقصد فایل", slug="default-file-target", owner=owner)
    WorkspaceMember.objects.create(workspace=other, member=owner, role=20)
    import_config(other, owner, config)
    destination = FieldDefinition.objects.get(workspace=other, code=field.code)
    assert destination.config["default"][0] != uploaded.data["id"]
    asset = FileAsset.objects.get(id=destination.config["default"][0])
    with asset.asset.open("rb") as stream:
        assert stream.read() == b"reusable default"
    import_config(other, owner, config)
    assert FileAsset.objects.filter(workspace=other, entity_type="CRM_FIELD_DEFAULT").count() == 1


def test_binding_defaults_are_typed_and_invalid_draft_is_rejected(engine):
    workspace, _, client = engine
    entity, form = make_entity(
        engine,
        [
            {"code": "range_default", "title": "بازه", "kind": "date_range"},
            {"code": "boolean_default", "title": "گزینه", "kind": "boolean", "config": {"default": True}},
            {"code": "tags_default", "title": "تگ", "kind": "tags"},
        ],
    )
    form.draft["sections"][0]["bindings"][1]["default"] = {"start": "2026-10-01", "end": "2026-10-05"}
    form.draft["sections"][0]["bindings"][2]["default"] = False
    form.draft["sections"][0]["bindings"][3]["default"] = ["الف", "ب"]
    saved = client.put(endpoint(workspace, "forms/sample/"), {"draft": form.draft, "revision": 1}, format="json")
    assert saved.status_code == 200, saved.data
    assert (
        client.post(
            endpoint(workspace, "forms/sample/"), {"revision": saved.data["revision"]}, format="json"
        ).status_code
        == 200
    )

    record = create_record(engine, entity, {"title": "پیش‌فرض‌ها"})
    assert record["values"]["range_default"] == {"start": "2026-10-01", "end": "2026-10-05"}
    assert record["values"]["boolean_default"] is False
    assert record["values"]["tags_default"] == ["الف", "ب"]
    assert FieldDefinition.objects.get(workspace=workspace, code="boolean_default").config["default"] is True
    draft = client.get(endpoint(workspace, "forms/sample/")).data
    draft["draft"]["sections"][0]["bindings"][1]["default"] = "invalid range"
    assert (
        client.put(
            endpoint(workspace, "forms/sample/"),
            {"draft": draft["draft"], "revision": draft["revision"]},
            format="json",
        ).status_code
        == 400
    )


@pytest.mark.parametrize("code", ["contains", "exact", "nested__key"])
def test_field_codes_cannot_be_interpreted_as_json_query_operators(engine, code):
    workspace, _, client = engine
    response = client.post(
        endpoint(workspace, "definitions/fields/"),
        {"code": code, "title": "کد رزرو شده", "kind": "short_text"},
        format="json",
    )
    assert response.status_code == 400, response.data
